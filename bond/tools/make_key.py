"""Sinh Key_YYYYMMDD.xlsx tu File 02 — day du 5 trang cua bao cao.

Nguon:
    Linked (1)   khoi 2  (52 chi tieu tuan thu han muc)
    Linked       khoi 3.x / 4 / 6.x / 7.x
    Chart data   19 chuoi time-series
    Report       nhan dinh + ghi chu
    Run Tool     truc ngay

Moi dong deu co khoa neo rieng, nen vi tri dong trong Excel khong con y nghia:
HTML tra cuu bang khoa, khong bang toa do.

    python3 make_key.py <File02.xlsm> [thu_muc_dich]
"""

import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="7B2D3B")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
EDIT_FONT = Font(color="0070C0")

BOOKS = [
    ("TB_MSB", "2.1.", "TRADING BOOK NỘI BỘ", 5, 30),
    ("BB_MSB", "2.2.", "BANKING BOOK NỘI BỘ", 32, 53),
    ("TB_SBV", "2.3.", "TRADING BOOK SBV", 55, 60),
    ("BB_SBV", "2.4.", "BANKING BOOK SBV", 62, 67),
    ("OTHER", "2.5.", "KHÁC", 69, 69),
    ("FIBOND", "2.6.", "FI Bond & CD", 71, 74),
]
S2COL = dict(stt=1, cat=2, label=3, today=4, dtd=5, yest=6, lm=7, lq=8, ly=9,
             limit=10, used=11, light=12)

POS = [
    ("TB", 14, 65, 75),
    ("BB", 28, 80, 90),
]
POS_FIELDS = ["Tenor", "Face", "FaceYest", "FaceLM", "DtD", "MtD",
              "PV01", "Itd", "ItdYest", "ItdLM", "ItdMtD", "ItdYtD", "Daily"]

CURVES = [
    ("YIELD", 41, 80, 89, ["Tenor", "Bid", "Ask", "Mid", "Spread", "", "DtD", "MtD", "YtD"]),
    ("REPO", 95, 171, 178, ["Tenor", "Today", "Yest", "LM", "", "", "DtD", "MtD"]),
]

GRIDS = [
    ("MIX", 147, 238, 249, 7, "3.3 Cơ cấu theo tổ chức phát hành",
     ["Tenor", "TB_TPCP", "TB_TPCPBL", "TB_TPCQDP", "BB_TPCP", "BB_TPCPBL", "BB_TPCQDP"]),
    ("HOLD", 164, 274, 284, 5, "3.4 Cơ cấu theo thời gian nắm giữ",
     ["Bucket", "TB_Pos", "TB_Itd", "BB_Pos", "BB_Itd"]),
    ("PNL", 154, 253, 259, 4, "3.5 Unrealized & Realized PnL",
     ["Label", "YtD", "MtD", "DtD"]),
    ("CAPITAL", 158, 262, 270, 6, "3.8 Mức độ sử dụng vốn",
     ["Label", "Today", "Yest", "DtD", "LM", "MtD"]),
    ("BS", 103, 184, 189, 14, "3.6 Ghi nhận PnL theo lớp bảng cân đối",
     ["Label", "On_MtM", "On_Face", "On_PTCK", "On_Diff",
      "Off_MtM", "Off_Face", "Off_PTCK", "Off_Diff",
      "Out_MtM", "Out_Face", "Out_PTCK", "Out_Diff", "Extra"]),
    ("FIONBS", 169, 288, 294, 5, "7.1 FI Bond trên bảng cân đối",
     ["Label", "OnBS", "OffBS", "OutBS", "Total"]),
]

SCENS = [
    ("TB", "recent", 51, 95, 105, 92),
    ("TB", "var", 51, 110, 120, 107),
    ("BB", "recent", 51, 140, 150, 137),
    ("BB", "var", 51, 125, 135, 122),
]

TEXTS = [
    ("txt.market", "B4", "Thông tin thị trường"),
    ("txt.assessment", "H4", "Tuân thủ hạn mức & đánh giá rủi ro"),
    ("txt.noteTB", "Q10", "Ghi chú Trading Book"),
    ("txt.noteTPCP", "Q59", "Ghi chú tỷ lệ đầu tư TPCP"),
    ("txt.noteFI", "Q66", "Ghi chú cơ cấu FI Bond"),
    ("txt.noteFV", "B110", "Ghi chú Fair value"),
    ("txt.itdTB", "B84", "Lỗ MtM Trading theo kỳ hạn"),
    ("txt.itdBB", "L84", "Lỗ MtM Banking theo kỳ hạn"),
    ("txt.note10d", "B166", "Kịch bản 10 ngày · Trading"),
    ("txt.noteVar", "B184", "Kịch bản VaR · Trading"),
    ("txt.noteBB10d", "B203", "Kịch bản 1 tháng · Banking"),
    ("txt.noteBBVar", "B221", "Kịch bản VaR · Banking"),
    ("txt.noteRealized", "M95", "Ghi chú Realized PnL"),
    ("txt.noteScenario", "M110", "Ghi chú kịch bản PnL"),
    ("txt.noteSign", "B120", "Quy ước dấu"),
    ("txt.noteRating", "B254", "Ghi chú xếp hạng FI Bond"),
    ("txt.noteVira", "N223", "Ghi chú VIRA"),
]


def slug(text):
    s = unicodedata.normalize("NFD", str(text))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D")
    s = re.sub(r"[^0-9a-zA-Z]+", " ", s).strip()
    parts = s.split()
    if not parts:
        return "unnamed"
    return parts[0].lower() + "".join(p.capitalize() for p in parts[1:])


def vndate(v):
    return v.strftime("%d/%m/%Y") if isinstance(v, (date, datetime)) else v


def clean(v):
    if isinstance(v, str):
        return re.sub(r"\s+", " ", v).strip()
    if isinstance(v, (date, datetime)):
        return vndate(v)
    return v


def rowvals(ws, r, c0, n):
    return [clean(ws.cell(r, c0 + i).value) for i in range(n)]


def build(src, outdir):
    wb = openpyxl.load_workbook(src, data_only=True)
    s2, lk, cd, rp, rt = (wb["Linked (1)"], wb["Linked"], wb["Chart data"],
                          wb["Report"], wb["Run Tool"])
    out = openpyxl.Workbook()
    counts = {}

    dates = {}
    for k, cell in (("today", "B2"), ("yest", "B3"), ("lm", "B4"), ("ly", "B7")):
        dates[k] = vndate(rt[cell].value)
    dates["lq"] = vndate(s2["H3"].value)

    meta = out.active
    meta.title = "META"
    meta.append(["Key", "Value", "Ghi chu"])
    for row in [
        ["schema", "bond.key.v2", "Phien ban cau truc file key"],
        ["asOf", dates["today"], "Ngay bao cao"],
        ["dateYest", dates["yest"], "Cot Yesterday"],
        ["dateLastMonth", dates["lm"], "Cot Last month"],
        ["dateLastQuarter", dates["lq"], "Cot Last Quarter"],
        ["dateLastYear", dates["ly"], "Cot Last Year"],
        ["rptTitle", "Báo cáo rủi ro thị trường", "Tieu de trang"],
        ["rptSubtitle", "Báo cáo Desk Bond", "Phu de"],
        ["unitNote", "Đơn vị: tỷ VND, trừ khi ghi khác. Giá trị âm là lỗ.", ""],
        ["generatedAt", datetime.now().strftime("%d/%m/%Y %H:%M"), "Thoi diem sinh file"],
    ]:
        meta.append(row)

    dsh = out.create_sheet("DATA_S2")
    dsh.append(["KeyID", "Book", "Code", "Section", "STT", "Nhom", "ChiTieu", "Sub",
                "Today", "DtD", "Yesterday", "LastMonth", "LastQuarter", "LastYear",
                "Limit", "Used", "Light"])
    seen = {}
    n = 0
    for key, code, title, r0, r1 in BOOKS:
        for r in range(r0, r1 + 1):
            label = clean(s2.cell(r, S2COL["label"]).value)
            if not label:
                continue
            base = slug(label)
            seen[key + base] = seen.get(key + base, 0) + 1
            if seen[key + base] > 1:
                base = "%s%d" % (base, seen[key + base])
            dsh.append(["S2.%s.%s" % (key, base), key, code, title,
                        s2.cell(r, S2COL["stt"]).value, clean(s2.cell(r, S2COL["cat"]).value),
                        label, 1 if re.match(r"^[ab]\)", label) else 0]
                       + [clean(s2.cell(r, S2COL[f]).value)
                          for f in ("today", "dtd", "yest", "lm", "lq", "ly", "limit", "used", "light")])
            n += 1
    counts["DATA_S2"] = n

    psh = out.create_sheet("POS")
    psh.append(["KeyID", "Book"] + POS_FIELDS)
    n = 0
    for book, c0, r0, r1 in POS:
        for r in range(r0, r1 + 1):
            v = rowvals(lk, r, c0, len(POS_FIELDS))
            if v[0] in (None, ""):
                continue
            psh.append(["POS.%s.%s" % (book, slug(v[0])), book] + v)
            n += 1
    counts["POS"] = n

    csh = out.create_sheet("CURVE")
    csh.append(["KeyID", "Type", "Tenor", "V1", "V2", "V3", "V4", "DtD", "MtD", "YtD"])
    n = 0
    for name, c0, r0, r1, fields in CURVES:
        for r in range(r0, r1 + 1):
            v = rowvals(lk, r, c0, 9)
            if v[0] in (None, ""):
                continue
            csh.append(["CURVE.%s.%s" % (name, slug(v[0])), name, v[0],
                        v[1], v[2], v[3], v[4], v[6], v[7], v[8]])
            n += 1
    counts["CURVE"] = n

    gsh = out.create_sheet("GRID")
    gsh.append(["KeyID", "Block", "BlockName", "Label"] + ["C%d" % i for i in range(1, 15)])
    n = 0
    for name, c0, r0, r1, width, title, fields in GRIDS:
        for r in range(r0, r1 + 1):
            v = rowvals(lk, r, c0, width)
            if v[0] in (None, ""):
                continue
            gsh.append(["GRID.%s.%s" % (name, slug(v[0])), name, title, v[0]]
                       + v[1:] + [None] * (14 - (width - 1)))
            n += 1
    counts["GRID"] = n

    ssh = out.create_sheet("SCEN")
    ssh.append(["KeyID", "Book", "Kind", "Scenario", "Sub", "Tenor", "PV01", "Itd",
                "YieldBps", "ItdChange"])
    n = 0
    for book, kind, c0, r0, r1, hdr in SCENS:
        names, subs = [], []
        for i in range(6):
            names.append(clean(lk.cell(hdr, c0 + 3 + i * 2).value))
            subs.append(clean(lk.cell(hdr + 2, c0 + 3 + i * 2).value))
        for r in range(r0, r1 + 1):
            tenor = clean(lk.cell(r, c0).value)
            if not tenor:
                continue
            pv01 = clean(lk.cell(r, c0 + 1).value)
            itd = clean(lk.cell(r, c0 + 2).value)
            for i in range(6):
                nm = names[i] or ("KB%d" % (i + 1))
                ssh.append(["SCEN.%s.%s.%s.%s" % (book, kind, slug(nm), slug(tenor)),
                            book, kind, nm, subs[i], tenor, pv01, itd,
                            clean(lk.cell(r, c0 + 3 + i * 2).value),
                            clean(lk.cell(r, c0 + 4 + i * 2).value)])
                n += 1
    counts["SCEN"] = n

    tsh = out.create_sheet("TEXT")
    tsh.append(["KeyID", "Mo ta", "Nguon", "Value"])
    for keyid, cell, desc in TEXTS:
        tsh.append([keyid, desc, "Report!" + cell, clean(rp[cell].value)])
    counts["TEXT"] = len(TEXTS)

    xsh = out.create_sheet("TS")
    xsh.append(["Series", "Field", "Date", "Value"])
    counts["TS"] = dump_ts(cd, xsh)

    style(meta, [18, 46, 40], 2)
    style(dsh, [30, 10, 7, 24, 6, 18, 32, 5] + [13] * 9)
    style(psh, [26, 7] + [12] * len(POS_FIELDS))
    style(csh, [26, 8, 9] + [11] * 7)
    style(gsh, [30, 10, 34, 34] + [12] * 14)
    style(ssh, [44, 7, 8, 26, 22, 8] + [11] * 4)
    style(tsh, [20, 34, 14, 95], 4, wrap=True)
    style(xsh, [16, 14, 13, 14])
    for sh in (dsh, psh, csh, gsh, ssh, xsh):
        sh.freeze_panes = "B2"

    stamp = re.sub(r"[^0-9]", "", str(dates["today"]))
    stamp = stamp[4:8] + stamp[2:4] + stamp[0:2] if len(stamp) == 8 else "unknown"
    dest = Path(outdir) / ("Key_%s.xlsx" % stamp)
    out.save(dest)
    return dest, counts


AXIS = ("Date", "Ngày", "Tháng", "Kỳ hạn", "STT")


def dump_ts(cd, sh):
    maxc = cd.max_column
    heads = [clean(cd.cell(1, c).value) for c in range(1, maxc + 1)]

    starts = []
    for i, h in enumerate(heads):
        if h and i + 1 < len(heads) and heads[i + 1] in AXIS:
            starts.append(i)

    n = 0
    for k, i in enumerate(starts):
        end = starts[k + 1] if k + 1 < len(starts) else len(heads)
        title, axis_col = heads[i], i + 2
        fields = [(c + 1, heads[c]) for c in range(i + 2, end) if heads[c]]
        series = slug(title)
        r = 2
        while r <= cd.max_row:
            d = cd.cell(r, axis_col).value
            if d in (None, ""):
                break
            for cc, f in fields:
                v = cd.cell(r, cc).value
                if v not in (None, ""):
                    sh.append([series, f, clean(d), v])
                    n += 1
            r += 1
    return n


def style(ws, widths, value_col=None, wrap=False):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if value_col:
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, value_col)
            cell.font = EDIT_FONT
            if wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path, counts = build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else ".")
    print(path)
    for k, v in counts.items():
        print("  %-10s %d dong" % (k, v))
