import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

exec(open('/home/user/HTML/bond/tools/ghep_cap_deal_fibond.py').read())

OUT = '/home/user/HTML/bond/Deal_FIBond_Repo_20260824.xlsx'
FONT = 'Arial'
HFONT = Font(name=FONT, size=10, bold=True)
CELL = Font(name=FONT, size=10)
DT = 'DD/MM/YYYY'
VND = '#,##0'
PC = '0.00%'

byid = {x['DEAL_ID']: x for x in R}
wb = openpyxl.Workbook()

COLS = ['Deal ban (S)', 'Deal mua (B)', 'Ma giay to', 'Product', 'Doi tac', 'To chuc phat hanh',
        'Menh gia (ty)', 'Chan 1 thanh toan', 'Chan 2 thanh toan', 'Ky han (ngay)',
        'Loai', 'Gross chan 1 (VND)', 'Gross chan 2 (VND)', 'Lai/lo (trieu)', '%/nam',
        'Yield S', 'Yield B', 'Ngay nhap S', 'Ngay nhap B', 'Lech ngay nhap']


def write_sheet(ws, rows):
    for i, t in enumerate(COLS, 1):
        ws.cell(row=1, column=i, value=t).font = HFONT
    for i, p in enumerate(sorted(rows, key=lambda p: p['d1']), 0):
        r = 2 + i
        isA = p['dirn'] == 'A'
        s, b = byid[p['sid']], byid[p['bid']]
        g1 = p['cash']
        g2 = p['cash'] - p['pnl'] if isA else p['cash'] + p['pnl']
        rate = ((-p['pnl'] if isA else p['pnl']) / p['cash'] * 365 / p['days']) if p['days'] and p['cash'] else None
        vals = [p['sid'], p['bid'], p['paper'], p['prod'], p['cpty'], s['CPTY_CODE_ISSUER'],
                p['face'] / 1e9, p['d1'], p['d2'], p['days'],
                p['loai'], g1, g2,
                p['pnl'] / 1e6, rate, s['YIELD'] / 100, b['YIELD'] / 100,
                p['scap'], p['bcap'], p['gap_cap']]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = CELL
        ws.cell(row=r, column=7).number_format = '#,##0.0'
        for ci in (8, 9, 18, 19):
            ws.cell(row=r, column=ci).number_format = DT
        for ci in (12, 13):
            ws.cell(row=r, column=ci).number_format = VND
        ws.cell(row=r, column=14).number_format = '#,##0.0'
        for ci in (15, 16, 17):
            ws.cell(row=r, column=ci).number_format = PC
    for i in range(1, len(COLS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.auto_filter.ref = f'A1:T{1 + len(rows)}'
    ws.freeze_panes = 'A2'


ws1 = wb.active
ws1.title = 'Repo cung ngay nhap'
write_sheet(ws1, repo_same_cap)

ws2 = wb.create_sheet('Repo khac ngay nhap BO')
write_sheet(ws2, repo_diff_cap)

wb.save(OUT)
print('saved', OUT)
print(f'  Sheet 1 "Repo cung ngay nhap"    : {len(repo_same_cap)} dong')
print(f'  Sheet 2 "Repo khac ngay nhap BO" : {len(repo_diff_cap)} dong')
print(f'  Tong Repo+Reverse Repo           : {len(matched)} dong')
