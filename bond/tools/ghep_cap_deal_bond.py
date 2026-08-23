#!/usr/bin/env python3
"""Ghép cặp repo từ export deal Bond (MSB_RP_DM) và bơm số vào file HTML báo cáo.

    python3 ghep_cap_deal_bond.py <deals.xlsx> [--html <bao_cao.html>] [--json D.json]

File export là danh sách PHẲNG từng chân giao dịch — không có cột nào nối chân bán
với chân mua lại của cùng một hợp đồng. Script này tái tạo lại quan hệ đó, rồi tính
lãi/lỗ theo cặp.

Chạy không có --html thì chỉ in thống kê + ghi JSON, không sửa file nào.
Có --html thì thay khối `const D={...};` trong file đó tại chỗ.

Ghi chú vận hành, xem thêm docs/bond-deal-pairing.md:
  * Cột GrossAmount có dòng là STRING dạng '  1012036120.0000000K' (đơn vị nghìn).
    Không parse riêng là lệch 1000 lần.
  * Ngày trong file là serial Excel, không phải datetime.
"""

import argparse
import datetime
import json
import re
import statistics
import sys
from collections import defaultdict

import openpyxl

# Cột được đọc, theo đúng thứ tự trong sheet export.
NCOLS = 25
SHEET = 'Sheet1'

# Ngày chốt sổ với TGĐ — mốc chia "đã đáo hạn" và "còn hiệu lực".
# Đổi khi kỳ báo cáo đổi.
AS_OF = datetime.datetime(2026, 8, 20)

# Cặp ghép CHÉO ĐỐI TÁC đã được nghiệp vụ xác nhận là repo (MSB đứng trung gian).
# Khoá theo (bid, sid). Cặp chéo nào KHÔNG có trong đây sẽ bị cảnh báo — phải hỏi
# lại nghiệp vụ trước khi dùng số, vì ghép chéo nhầm tạo ra cặp không có thật mà
# vẫn khớp khối lượng.
CROSS_OK = {
    (50120, '50127+50128'),   # mua KBNN (IB-P-VSD) -> bán PGBV-HO (OT-P-VSD)
}


# --------------------------------------------------------------------------- đọc
def xl2dt(v):
    """Serial Excel -> datetime. Gốc 1899-12-30 (không phải 1900-01-01)."""
    if isinstance(v, datetime.datetime):
        return v
    return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=v)


def parse_gross(v):
    """GrossAmount: đa số là số, nhưng một phần là string '  <số>K' — đơn vị nghìn."""
    if isinstance(v, str):
        return float(v.strip().rstrip('K')) * 1000.0
    return float(v)


def read_deals(path):
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    headers = [ws.cell(row=1, column=c).value for c in range(1, NCOLS + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        d = dict(zip(headers, [ws.cell(row=r, column=c).value for c in range(1, NCOLS + 1)]))
        if d.get('BondsDeals_Id') is None:
            continue
        d['GrossAmount'] = parse_gross(d['GrossAmount'])
        for k in ('TradeDate', 'SettlementDate', 'CaptureDate', 'IssueDate', 'MaturityDate'):
            d[k] = xl2dt(d[k])
        rows.append(d)
    return rows


# ------------------------------------------------------------------------- ghép
def emit(s, b, qty, split, bond, cpty, out):
    """Dựng một cặp từ chân bán s và chân mua b, với khối lượng qty (đơn vị TP).

    Nhóm do CHÂN NÀO THANH TOÁN TRƯỚC quyết định, không phải do S hay B:
      A = chân bán trước  -> MSB nhận tiền rồi trả lại  -> MSB ĐI VAY
      B = chân mua trước  -> MSB bỏ tiền rồi thu về     -> MSB CHO VAY

    Dấu thống nhất một chiều cho cả hai nhóm: cost = tiền chân B − tiền chân S.
    Nhóm A cost dương = trả lãi (chi phí). Nhóm B cost âm = thu lãi.
    Việc đổi dấu cho dễ đọc để phía trình bày lo.
    """
    dirn = 'A' if s['SettlementDate'] <= b['SettlementDate'] else 'B'
    days = abs((b['SettlementDate'] - s['SettlementDate']).days)
    s_cash = s['GrossAmount'] * (qty / s['Quantity'])
    b_cash = b['GrossAmount'] * (qty / b['Quantity'])
    cash_bn = (s_cash if dirn == 'A' else b_cash) / 1e9      # tiền chân đầu, tỷ VND
    cost_mn = (b_cash - s_cash) / 1e6                        # triệu VND
    rate = (cost_mn / (cash_bn * 1000) * 365 / days * 100) if days > 0 and cash_bn > 0 else 0.0
    out.append({
        'sid': s['BondsDeals_Id'], 'bid': b['BondsDeals_Id'], 'bond': bond, 'cpty': cpty,
        'qty': round(qty / 1000, 3),
        'sset': s['SettlementDate'].strftime('%Y-%m-%d'),
        'bset': b['SettlementDate'].strftime('%Y-%m-%d'), 'days': days,
        'cash': round(cash_bn, 5), 'cost': round(cost_mn, 4), 'rate': rate,
        'dy': (b['Yield'] - s['Yield']) * 100,               # điểm cơ bản
        'dirn': dirn,
        'tag': 'exact' if not split else
               ('same-cap' if s['CaptureDate'] == b['CaptureDate'] else 'cross-cap'),
        'sf': s['Folders_ShortName'], 'bf': b['Folders_ShortName'],
        'sy': s['Yield'], 'by': b['Yield'], 'cpn': s['CouponRate'],
        'mat': s['MaturityDate'].strftime('%Y-%m-%d'),
        'sclr': s['ClearingModes_ShortName'], 'bclr': b['ClearingModes_ShortName'],
        'scap': s['CaptureDate'].strftime('%Y-%m-%d'), 'bcap': b['CaptureDate'].strftime('%Y-%m-%d'),
        'std': s['TradeDate'].strftime('%Y-%m-%d'), 'btd': b['TradeDate'].strftime('%Y-%m-%d'),
    })


def match_same_cpty(rows):
    """Ghép trong phạm vi (mã TP, đối tác) — điều kiện của một hợp đồng repo thật.

    Trả về (danh sách cặp, danh sách chân còn dư).
    """
    groups = defaultdict(lambda: {'S': [], 'B': []})
    for r in rows:
        groups[(r['Bonds_ShortName'], r['Cpty_ShortName'])][r['DealType']].append(dict(r))

    pairs, leftover = [], []
    for (bond, cpty), g in groups.items():
        slegs = sorted(g['S'], key=lambda x: (x['SettlementDate'], x['BondsDeals_Id']))
        blegs = sorted(g['B'], key=lambda x: (x['SettlementDate'], x['BondsDeals_Id']))
        sq = [x['Quantity'] for x in slegs]
        bq = [x['Quantity'] for x in blegs]

        # Lượt 1 — khôi phục cặp theo dấu vết book lệnh: cùng CaptureDate, cùng
        # Quantity, ưu tiên deal-id gần nhau nhất. Hai chân của một hợp đồng được
        # nhập cùng lúc nên dính nhau về ngày nhập và số hiệu; tiêu chí này thắng
        # thứ tự thanh toán thuần túy khi nhiều chân trùng ngày.
        cands = sorted(
            (abs(s['BondsDeals_Id'] - b['BondsDeals_Id']), i, j)
            for i, s in enumerate(slegs) for j, b in enumerate(blegs)
            if s['CaptureDate'] == b['CaptureDate'] and s['Quantity'] == b['Quantity']
        )
        for _, i, j in cands:
            # chỉ ghép khi CẢ HAI chân còn nguyên: lượt này không tách chân
            if sq[i] != slegs[i]['Quantity'] or bq[j] != blegs[j]['Quantity']:
                continue
            if sq[i] <= 1e-6 or bq[j] <= 1e-6:
                continue
            qty, sq[i], bq[j] = sq[i], 0.0, 0.0
            emit(slegs[i], blegs[j], qty, False, bond, cpty, pairs)

        # Lượt 2 — FIFO phần dư theo ngày thanh toán, tách chân lớn nếu lệch khối lượng.
        si = bi = 0
        while si < len(slegs) and bi < len(blegs):
            if sq[si] <= 1e-6:
                si += 1
                continue
            if bq[bi] <= 1e-6:
                bi += 1
                continue
            s, b = slegs[si], blegs[bi]
            qty = min(sq[si], bq[bi])
            split = (s['Quantity'] != qty) or (b['Quantity'] != qty)
            sq[si] -= qty
            bq[bi] -= qty
            emit(s, b, qty, split, bond, cpty, pairs)

        for arr, q, side in ((slegs, sq, 'S'), (blegs, bq, 'B')):
            leftover += [{**x, 'side': side, 'rem': q[i]} for i, x in enumerate(arr) if q[i] > 1e-6]

    return pairs, leftover


def match_cross_cpty(leftover):
    """Nới điều kiện "cùng đối tác" cho phần dư: mua của bên này, bán cho bên kia.

    VẪN LÀ REPO — hai chân khác đối tác là do MSB đứng trung gian (đối tác làm việc
    với chính phủ). Trường hợp thực tế: mua của KBNN (liên ngân hàng) rồi bán cho
    PGBV-HO (OTC). Gắn tag 'est' để phía báo cáo nhận ra, không phải để loại ra.

    Quy tắc ghép chuẩn là cùng mã TP + cùng đối tác; lượt này là ngoại lệ chỉ dùng
    cho deal nghiệp vụ đã xác nhận. Cặp nào script tự ghép chéo mà chưa ai xác nhận
    thì phải hỏi lại nghiệp vụ trước khi dùng số — xem cảnh báo in ra ở cuối main().
    """
    by_bond = defaultdict(list)
    for x in leftover:
        by_bond[x['Bonds_ShortName']].append(x)

    cross, consumed = [], set()
    for bond, legs in by_bond.items():
        slegs = sorted((x for x in legs if x['side'] == 'S'), key=lambda x: x['SettlementDate'])
        blegs = sorted((x for x in legs if x['side'] == 'B'), key=lambda x: x['SettlementDate'])
        for b in blegs:
            if b['BondsDeals_Id'] in consumed:
                continue
            remaining, taken = b['rem'], []
            for s in slegs:
                if remaining <= 1e-6:
                    break
                if s['BondsDeals_Id'] in consumed or s['rem'] <= 1e-6:
                    continue
                t = min(remaining, s['rem'])
                taken.append((s, t))
                remaining -= t
                s['rem'] -= t
                if s['rem'] <= 1e-6:
                    consumed.add(s['BondsDeals_Id'])
            if not taken or remaining > 1e-6:
                continue

            consumed.add(b['BondsDeals_Id'])
            s0 = taken[0][0]
            total = sum(t for _, t in taken)
            s_cash = sum(s['GrossAmount'] * (t / s['Quantity']) for s, t in taken)
            b_cash = b['GrossAmount'] * (total / b['Quantity'])
            bset = b['SettlementDate']
            sset = max(s['SettlementDate'] for s, _ in taken)
            dirn = 'A' if bset >= sset else 'B'
            days = abs((sset - bset).days)
            cash_bn = (b_cash if dirn == 'B' else s_cash) / 1e9
            cost_mn = (b_cash - s_cash) / 1e6
            cross.append({
                'sid': '+'.join(str(s['BondsDeals_Id']) for s, _ in taken),
                'bid': b['BondsDeals_Id'], 'bond': bond,
                'cpty': f"{b['Cpty_ShortName']}→{s0['Cpty_ShortName']}",
                'qty': round(total / 1000, 3),
                'sset': sset.strftime('%Y-%m-%d'), 'bset': bset.strftime('%Y-%m-%d'), 'days': days,
                'cash': round(cash_bn, 5), 'cost': round(cost_mn, 4),
                'rate': (cost_mn / (cash_bn * 1000) * 365 / days * 100) if days > 0 and cash_bn > 0 else 0.0,
                'dy': 0.0, 'dirn': dirn, 'tag': 'est',
                'sf': s0['Folders_ShortName'], 'bf': b['Folders_ShortName'],
                'cptyB': b['Cpty_ShortName'], 'cptyS': s0['Cpty_ShortName'],
                'sy': s0['Yield'], 'by': b['Yield'], 'cpn': b['CouponRate'],
                'mat': b['MaturityDate'].strftime('%Y-%m-%d'),
                'sclr': s0['ClearingModes_ShortName'], 'bclr': b['ClearingModes_ShortName'],
                'scap': s0['CaptureDate'].strftime('%Y-%m-%d'), 'bcap': b['CaptureDate'].strftime('%Y-%m-%d'),
                'std': s0['TradeDate'].strftime('%Y-%m-%d'), 'btd': b['TradeDate'].strftime('%Y-%m-%d'),
            })

    still = [x for x in leftover if x['BondsDeals_Id'] not in consumed]
    return cross, still


# ------------------------------------------------------------------- tổng hợp
def dt(s):
    return datetime.datetime.strptime(s, '%Y-%m-%d')


def close_dt(p):
    """Ngày thanh toán chân sau — thời điểm lãi/lỗ thực sự được ghi nhận."""
    return max(dt(p['sset']), dt(p['bset']))


def is_anom(p):
    """Cờ định giá lại, CÓ TÍNH ĐẾN KỲ HẠN.

    Repo ngắn thì hai chân phải cùng một mức yield thỏa thuận; lệch ≥ 10bp nghĩa là
    chân mua lại bị mark theo thị trường — lãi/lỗ khi đó là rủi ro giá chứ không phải
    chi phí vốn. Kỳ hạn dài thì hai yield lệch nhau là chuyện bình thường, cùng ngưỡng
    đó không nói lên gì và không được loại.
    """
    return abs(p['dy']) >= 10 and p['days'] <= 30


def wrate(ps, sign):
    """Lãi suất ngụ ý %/năm, trọng số theo tiền × ngày."""
    den = sum(p['cash'] * p['days'] / 365 for p in ps)
    return (sum(sign * p['cost'] / 1e3 for p in ps) / den * 100) if den else 0.0


def blk(ps, sign):
    cl = [p for p in ps if not is_anom(p)]
    done = [p for p in ps if close_dt(p) <= AS_OF]
    op = [p for p in ps if close_dt(p) > AS_OF]
    return {
        'nDone': len(done), 'pnlDone': round(sum(sign * p['cost'] for p in done), 2),
        'cashDone': round(sum(p['cash'] for p in done), 3),
        'nOpen': len(op), 'pnlOpen': round(sum(sign * p['cost'] for p in op), 2),
        'cashOpen': round(sum(p['cash'] for p in op), 3),
        'n': len(ps), 'cash': round(sum(p['cash'] for p in ps), 3),
        'pnl': round(sum(sign * p['cost'] for p in ps), 2), 'rate': wrate(ps, sign),
        'nX': len(cl), 'pnlX': round(sum(sign * p['cost'] for p in cl), 2), 'rateX': wrate(cl, sign),
        'med': statistics.median([p['days'] for p in ps]) if ps else 0,
        'qty': round(sum(p['qty'] for p in ps), 1),
    }


def group_by(ps, sign, key, extra=None, top=None):
    out = []
    g = defaultdict(list)
    for p in ps:
        g[key(p)].append(p)
    for k, v in g.items():
        b = blk(v, sign)
        b['k'] = k
        if extra:
            b.update(extra(v))
        out.append(b)
    out.sort(key=lambda x: -x['cash'])
    return out[:top] if top else out


def outstanding(ps):
    """Dư nợ còn hiệu lực tại từng ngày lịch, tỷ VND."""
    ev = defaultdict(float)
    for p in ps:
        a, z = sorted([dt(p['sset']), dt(p['bset'])])
        ev[a] += p['cash']
        ev[z] -= p['cash']
    if not ev:
        return []
    cur, out, d, end = 0.0, [], min(ev), max(ev)
    while d <= end:
        cur += ev.get(d, 0.0)
        out.append([d.strftime('%Y-%m-%d'), round(cur, 1)])
        d += datetime.timedelta(days=1)
    return out


def build_D(pairs, unmatched, n_deals, d_from, d_to):
    A = [p for p in pairs if p['dirn'] == 'A']
    B = [p for p in pairs if p['dirn'] == 'B']

    def tenor(ps, sign):
        out = []
        for k, lo, hi in [('1 ngày', 1, 1), ('2–7 ngày', 2, 7), ('8–30 ngày', 8, 30),
                          ('31–90 ngày', 31, 90), ('> 90 ngày', 91, 10 ** 9)]:
            v = [p for p in ps if lo <= p['days'] <= hi]
            if v:
                b = blk(v, sign)
                b['k'] = k
                out.append(b)
        return out

    def month(ps, sign, dkey):
        g = defaultdict(list)
        for p in ps:
            g[dt(p[dkey]).strftime('%m/%Y')].append(p)
        out = []
        for m in sorted(g, key=lambda x: (x[3:], x[:2])):
            b = blk(g[m], sign)
            b['m'] = m
            out.append(b)
        return out

    def cpty(ps, sign):
        r = group_by(ps, sign, lambda p: p['cpty'],
                     lambda v: {'minD': min(x['days'] for x in v), 'maxD': max(x['days'] for x in v)})
        for x in r:
            x['cpty'] = x.pop('k')
        return r

    def bond(ps, sign):
        r = group_by(ps, sign, lambda p: p['bond'],
                     lambda v: {'mat': v[0]['mat'], 'cpn': v[0]['cpn']}, top=14)
        for x in r:
            x['b'] = x.pop('k')
        return r

    def folder(ps, sign):
        r = group_by(ps, sign, lambda p: p['sf'] if p['dirn'] == 'A' else p['bf'])
        for x in r:
            x['f'] = x.pop('k')
        return r

    # anom = cần soát lại · noise = %/năm phóng đại do làm tròn giá trên deal qua đêm
    anom, noise = [], []
    for p in pairs:
        q = dict(p)
        q['r'] = (1 if p['dirn'] == 'A' else -1) * p['rate']
        if is_anom(p):
            q['why'] = 'Chân mua lại được định giá lại (yield lệch ≥ 10bp, kỳ hạn ≤ 30 ngày)'
            anom.append(q)
        elif p['days'] >= 3 and abs(q['r']) > 15:
            q['why'] = 'Lãi suất ngụ ý lệch xa mặt bằng trên kỳ hạn ≥ 3 ngày'
            anom.append(q)
        elif p['days'] <= 2 and abs(q['r']) > 8:
            q['why'] = 'Làm tròn giá trên deal qua đêm'
            noise.append(q)
    anom.sort(key=lambda x: -abs(x['cost']))
    noise.sort(key=lambda x: -abs(x['cost']))

    return {
        'kpi': {'nDeals': n_deals, 'nPairs': len(pairs), 'nUn': len(unmatched),
                'A': blk(A, 1), 'B': blk(B, -1), 'dFrom': d_from, 'dTo': d_to},
        'cptyA': cpty(A, 1), 'cptyB': cpty(B, -1),
        'tenorA': tenor(A, 1), 'tenorB': tenor(B, -1),
        'monthA': month(A, 1, 'sset'), 'monthB': month(B, -1, 'bset'),
        'bondA': bond(A, 1), 'bondB': bond(B, -1),
        'foldA': folder(A, 1), 'foldB': folder(B, -1),
        'outstA': outstanding(A), 'outstB': outstanding(B),
        'histA': [round(p['rate'], 3) for p in A if p['days'] > 0],
        'histB': [round(-p['rate'], 3) for p in B if p['days'] > 0],
        'anom': anom, 'noise': noise,
        'fmis': [p for p in pairs if p['sf'] != p['bf']],
        'cmis': [p for p in pairs if p['sclr'] != p['bclr']],
        'zero': [p for p in pairs if p['days'] == 0],
        'pairs': pairs, 'unmatched': unmatched,
    }


# ------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx', help='File export deal từ MSB_RP_DM')
    ap.add_argument('--html', help='File HTML báo cáo, thay khối const D tại chỗ')
    ap.add_argument('--json', help='Ghi D ra file JSON')
    a = ap.parse_args()

    rows = read_deals(a.xlsx)
    caps = [r['CaptureDate'] for r in rows]
    d_from, d_to = min(caps).strftime('%Y-%m-%d'), max(caps).strftime('%Y-%m-%d')
    print(f'{len(rows)} chân deal · CaptureDate {d_from} → {d_to}')

    pairs, leftover = match_same_cpty(rows)
    print(f'  lượt 1+2, cùng đối tác : {len(pairs)} cặp, dư {len(leftover)} chân')
    cross, unmatched_legs = match_cross_cpty(leftover)
    print(f'  ghép chéo đối tác      : {len(cross)} cặp')
    chua_xn = []
    for c in cross:
        ok = (c['bid'], str(c['sid'])) in CROSS_OK
        print(f'      {c["bid"]} ↔ {c["sid"]}  {c["bond"]}  {c["cpty"]}  {c["days"]}n'
              f'  {"[đã xác nhận]" if ok else "[CHƯA XÁC NHẬN]"}')
        if not ok:
            chua_xn.append(c)
    if chua_xn:
        print()
        print(f'  !! {len(chua_xn)} cặp ghép chéo đối tác CHƯA được nghiệp vụ xác nhận.')
        print('     Quy tắc ghép là cùng mã TP + cùng đối tác; ghép chéo chỉ áp cho deal')
        print('     đã chỉ đích danh. HỎI LẠI NGHIỆP VỤ trước khi dùng số của các cặp này,')
        print('     rồi thêm vào CROSS_OK ở đầu file.')

    pairs += cross
    unmatched = [{'id': x['BondsDeals_Id'], 'side': x['side'], 'bond': x['Bonds_ShortName'],
                  'cpty': x['Cpty_ShortName'], 'set': x['SettlementDate'].strftime('%Y-%m-%d'),
                  'qty': round(x['rem'] / 1000, 3)} for x in unmatched_legs]
    if unmatched:
        print(f'  CHÂN KHÔNG GHÉP ĐƯỢC   : {len(unmatched)} — cần rà lại input')
        for x in unmatched:
            print('     ', x)
    else:
        print('  chân không ghép được   : 0 — toàn bộ khối lượng đã offset')

    D = build_D(pairs, unmatched, len(rows), d_from, d_to)
    A, B = D['kpi']['A'], D['kpi']['B']
    print(f"\nNhóm A · đi vay  : {A['n']:3d} cặp · {A['cash']:>10,.0f} tỷ · "
          f"trung vị {A['med']:>5.0f} ngày · chi phí  {A['rateX']:6.2f}%/năm")
    print(f"Nhóm B · cho vay : {B['n']:3d} cặp · {B['cash']:>10,.0f} tỷ · "
          f"trung vị {B['med']:>5.0f} ngày · lợi suất {B['rateX']:6.2f}%/năm")
    print(f"Lãi/lỗ đã đáo hạn tính đến {AS_OF:%d/%m/%Y}: "
          f"{(B['pnlDone'] - A['pnlDone']) / 1000:,.1f} tỷ "
          f"· còn hiệu lực theo hợp đồng {(B['pnlOpen'] - A['pnlOpen']) / 1000:,.1f} tỷ")
    print(f"Cần soát lại: {len(D['anom'])} cặp · nhiễu làm tròn qua đêm: {len(D['noise'])} cặp "
          f"· lệch folder: {len(D['fmis'])} cặp")

    blob = json.dumps(D, ensure_ascii=False, separators=(',', ':'))
    if a.json:
        open(a.json, 'w').write(blob)
        print(f'\n→ {a.json}  ({len(blob):,} bytes)')

    if a.html:
        src = open(a.html, encoding='utf-8').read()
        new, n = re.subn(r'const D=\{.*?\};\n', 'const D=' + blob + ';\n', src, count=1, flags=re.S)
        if not n:
            sys.exit(f'LỖI: không tìm thấy khối `const D={{...}};` trong {a.html}')
        open(a.html, 'w', encoding='utf-8').write(new)
        print(f'→ {a.html}  đã thay khối const D')


if __name__ == '__main__':
    main()
